# -*- coding: utf-8 -*-
"""把 sync.cjs 导出的 files_list.json 生成 Excel 目录：AI问答文档目录.xlsx"""
import json
import os

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "files_list.json")
OUT = os.path.join(os.path.dirname(BASE), "AI问答文档目录.xlsx")

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open(SRC, encoding="utf-8") as f:
    items = json.load(f)

items.sort(key=lambda x: (x.get("mtime") or ""), reverse=True)

wb = Workbook()
ws = wb.active
ws.title = "文档目录"

headers = ["序号", "文件名", "库内标题", "大小KB", "原位置", "文件生成时间", "最后修改时间", "上传ima时间", "状态"]
widths = [6, 52, 40, 9, 70, 18, 18, 18, 16]
thin = Side(style="thin", color="D0D0D0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="1F4E79")
head_font = Font(color="FFFFFF", bold=True, size=11)

for c, (h, w) in enumerate(zip(headers, widths), 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.fill = head_fill
    cell.font = head_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 22

status_icon = {
    "内容一致": "✅ 内容一致",
    "大小一致": "✅ 大小一致",
    "本机已上传": "✅ 本机已上传",
    "已在知识库": "✅ 已在知识库",
    "已上传": "✅ 已上传",
    "内容不同·已补传": "🆕 内容不同·已补传",
    "内容不同待补传": "🆕 内容不同·待补传",
    "待上传": "🆕 待上传",
    "待校验(下轮)": "⏸️ 待下轮校验",
    "同名跳过": "🔁 同名已在库",
    "超10MB限制": "🚫 超10MB",
}
link_font = Font(color="0563C1", underline="single")

for i, it in enumerate(items, 1):
    folder = os.path.dirname(it.get("path", ""))
    full_path = it.get("path", "")
    name = it.get("name", "")
    title = it.get("title") or name
    row = [
        i,
        name,
        "" if title == name else title,
        round((it.get("size") or 0) / 1024, 1),
        folder,
        it.get("birth") or it.get("mtime", ""),
        it.get("mtime", ""),
        it.get("uploadedAt") or "",
        status_icon.get(it.get("status", ""), it.get("status", "")),
    ]
    r = i + 1
    for c, v in enumerate(row, 1):
        cell = ws.cell(row=r, column=c)
        # 文件名列 → HYPERLINK 公式，点击直接打开本地 md
        if c == 2 and full_path:
            cell.value = f'=HYPERLINK("{full_path}","{name}")'
            cell.font = link_font
        else:
            cell.value = v
        if c == 3 and v:
            cell.font = Font(color="C00000")
        cell.border = border
        cell.alignment = Alignment(vertical="center", horizontal="center" if c in (1, 4, 6, 7, 8) else "left")
    if i % 2 == 0:
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="F2F7FB")

# KPI 汇总 sheet
ws2 = wb.create_sheet("KPI汇总")
uploaded = sum(1 for x in items if x.get("uploadedAt"))
pending = sum(1 for x in items if x.get("status") in ("待上传", "内容不同待补传"))
oversize = sum(1 for x in items if x.get("status") == "超10MB限制")
renamed = sum(1 for x in items if (x.get("title") or x.get("name")) != x.get("name"))
total_kb = sum((x.get("size") or 0) for x in items) / 1024
kpis = [
    ("指标", "数值"),
    ("文档总数", len(items)),
    ("已上传ima", uploaded),
    ("待上传", pending),
    ("库里同名但内容不同·加后缀另存", renamed),
    ("超10MB跳过", oversize),
    ("总大小(MB)", round(total_kb / 1024, 1)),
    ("最新文档时间", max((x.get("mtime") or "") for x in items) if items else ""),
]
for r, (k, v) in enumerate(kpis, 1):
    a = ws2.cell(row=r, column=1, value=k)
    b = ws2.cell(row=r, column=2, value=v)
    if r == 1:
        a.font = b.font = Font(bold=True, color="FFFFFF")
        a.fill = b.fill = head_fill
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 24

try:
    wb.save(OUT)
    saved = OUT
except PermissionError:
    saved = OUT.replace(".xlsx", "_更新.xlsx")
    wb.save(saved)
print(f"OK {saved} 共{len(items)}行 已上传{uploaded}")
